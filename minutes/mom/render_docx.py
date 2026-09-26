"""The same parsed body and meeting details -> an editable DOCX that matches
the PDF: logo, teal headings, decisions and action tables, the confirmation
box, signatures, and the AI notice in the footer and document properties."""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .latexcheck import unescape
from .render_pdf import DOC_WORD, TEMPLATE
from .schemas import Meeting, format_date, fresh

TEAL, SLATE, CHARCOAL, GREY = RGBColor(0x00, 0x82, 0x86), RGBColor(0x51, 0x59, 0x63), RGBColor(0x40, 0x3E, 0x3D), RGBColor(0xB3, 0xB7, 0xBA)
HEAD, BODY = "Montserrat", "PT Serif"


def strings(lang: str) -> dict:
    """Fixed labels, read from the LaTeX class so both formats say the same thing."""
    cls = (TEMPLATE / "medpark-mom.cls").read_text(encoding="utf-8")
    out = {}
    for m in re.finditer(r"\\momdefstring\s*\{(\w+)\}\s*\{((?:[^{}]|\{[^{}]*\})*)\}\s*\{((?:[^{}]|\{[^{}]*\})*)\}\s*\{((?:[^{}]|\{[^{}]*\})*)\}", cls):
        key, ro, ru, en = m.groups()
        out[key] = {"ro": ro, "ru": ru, "en": en}[lang]
    return out


def render(meeting: Meeting, blocks, lang: str, model: str, verified: str, out_path: Path) -> Path:
    s = strings(lang)
    s = {k: v.replace("\\momvalue{model}", model) for k, v in s.items()}
    doc = Document()
    _page(doc, s, meeting, verified)
    _styles(doc)
    _properties(doc, meeting, lang, model)

    _para(doc, f"{s['doc']} {s['no']} {meeting.number or '1'}", HEAD, 19, TEAL, bold=True, after=2)
    _para(doc, meeting.title(lang), HEAD, 12.5, CHARCOAL, after=8)
    meta = [(s["date"], format_date(meeting.date, lang)),
            (s["time"], f"{meeting.start}\u2013{meeting.end}" if meeting.start and meeting.end else meeting.start),
            (s["place"], meeting.place), (s["chair"], meeting.chair), (s["secretary"], meeting.secretary)]
    table = doc.add_table(rows=0, cols=2)
    for label, value in meta:
        row = table.add_row().cells
        _cell(row[0], label, HEAD, 9, SLATE)
        _cell(row[1], value, BODY, 10.5, CHARCOAL)

    _widths(table, [Mm(40), Mm(128)])
    for key, people in (("present", meeting.attendees), ("apologies", meeting.apologies)):
        if people:
            head = s[key] + (f"  ({meeting.quorum})" if key == "present" and meeting.quorum else "")
            _para(doc, head, HEAD, 10.5, TEAL, before=8, after=2)
            for p in people:
                par = doc.add_paragraph(style="Normal")
                _run(par, p.get("name", ""), BODY, 10.5, CHARCOAL)
                if p.get("role"):
                    _run(par, f", {p['role']}", BODY, 10.5, SLATE)
                par.paragraph_format.space_after = Pt(0)

    decisions, actions, confirm, n_topic = [], [], [], 0
    for b in blocks:
        a = {k: unescape(v) for k, v in b.args.items()}
        if b.kind == "summary":
            _para(doc, s["summary"], HEAD, 10.5, TEAL, before=8, after=2)
            _para(doc, a["text"], BODY, 10.5, CHARCOAL, after=4)
        elif b.kind == "begin" and a.get("env") == "agenda":
            _para(doc, s["agenda"], HEAD, 10.5, TEAL, before=8, after=2)
        elif b.kind == "agendaitem":
            doc.add_paragraph(a["title"], style="List Number")
        elif b.kind == "topic":
            n_topic += 1
            _para(doc, f"{n_topic}.  {a['title']}", HEAD, 12.5, TEAL, bold=True, before=12, after=3)
        elif b.kind in ("presented", "noted"):
            _para(doc, a["text"], BODY, 10.5, CHARCOAL, after=4)
        elif b.kind == "decision":
            decisions.append(a)
            _item(doc, f"{s['decision']} {len(decisions)}", a["text"], f"({a['vote']})" if a["vote"] else "")
        elif b.kind == "action":
            actions.append(a)
            detail = f"{s['owner']}: {a['owner']}" + (f"    {s['deadline']}: {a['deadline']}" if a["deadline"] else "")
            _item(doc, f"{s['task']} {len(actions)}", a["text"], detail, newline=True)
        elif b.kind == "needsconfirmation":
            confirm.append(a["text"])
        elif b.kind == "nextmeeting":
            par = _para(doc, "", BODY, 10.5, CHARCOAL, before=6)
            _run(par, f"{s['next']}: ", HEAD, 10.5, TEAL)
            _run(par, a["text"], BODY, 10.5, CHARCOAL)

    if decisions:
        _para(doc, s["decisions"], HEAD, 12.5, TEAL, bold=True, before=14, after=4)
        _table(doc, ["#", s["decision"]], [[str(i), d["text"] + (f"  ({d['vote']})" if d["vote"] else "")]
                                          for i, d in enumerate(decisions, 1)], [Mm(9), Mm(159)])
    if actions:
        _para(doc, s["actions"], HEAD, 12.5, TEAL, bold=True, before=14, after=4)
        _table(doc, ["#", s["task"], s["owner"], s["deadline"]],
               [[str(i), x["text"], x["owner"], x["deadline"] or s["notset"]] for i, x in enumerate(actions, 1)],
               [Mm(9), Mm(95), Mm(38), Mm(26)])
    if confirm:
        box = doc.add_table(rows=1, cols=1)
        cell = box.rows[0].cells[0]
        _widths(box, [Mm(168)])
        _shade(cell, "FCF8DC")
        _cell(cell, s["confirm"], HEAD, 10, CHARCOAL)
        _cell(cell, s["confirmhint"], BODY, 9, CHARCOAL, new=True)
        for text in confirm:
            _cell(cell, f"\u2022  {text}", BODY, 10, CHARCOAL, new=True)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    sig = doc.add_table(rows=3, cols=2)
    _widths(sig, [Mm(84), Mm(84)])
    for col, (label, name) in enumerate(((s["chair"], meeting.chair), (s["secretary"], meeting.secretary))):
        _cell(sig.rows[0].cells[col], label, HEAD, 9, SLATE)
        _cell(sig.rows[1].cells[col], "\n______________________________", BODY, 10, CHARCOAL)
        _cell(sig.rows[2].cells[col], name, BODY, 10, CHARCOAL)
    doc.paragraphs[-1].paragraph_format.space_before = Pt(18)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(fresh(out_path))
    out_path.chmod(0o600)
    return out_path


# ---------------------------------------------------------------- helpers
def _page(doc, s, meeting, verified):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(21)
    sec.top_margin, sec.bottom_margin = Mm(30), Mm(24)
    head = sec.header.paragraphs[0]
    head.add_run().add_picture(str(TEMPLATE / "logo.png"), width=Mm(30))
    right = sec.header.add_paragraph()
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(right, f"{s['doc']} {s['no']} {meeting.number or '1'}    {format_date(meeting.date, _lang_of(s))}", HEAD, 8.5, SLATE)
    foot = sec.footer.paragraphs[0]
    _run(foot, f"{s['aifooter']} {verified} {s['verified']}.\nMedpark | Spital Internațional, Sf. Andrei Doga 24, "
               f"MD-2024 Chișinău, +373 22 40 00 40. {s['confidential']}", HEAD, 7, SLATE)
    pages = sec.footer.add_paragraph()
    pages.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(pages, f"{s['page']} ", HEAD, 7, SLATE)
    _field(pages, "PAGE")
    _run(pages, f" {s['of']} ", HEAD, 7, SLATE)
    _field(pages, "NUMPAGES")


def _lang_of(s):
    return {"PROCES-VERBAL": "ro", "ПРОТОКОЛ": "ru"}.get(s["doc"], "en")


def _styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size, normal.font.color.rgb = BODY, Pt(10.5), CHARCOAL
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.1


def _properties(doc, meeting, lang, model):
    p = doc.core_properties
    p.title = f"{DOC_WORD[lang]} {meeting.title(lang)}"
    p.author = f"Secure MOM (AI-generated, local model {model})"
    p.subject = "AI-generated minutes of meeting. Drafted locally by an AI system and reviewed by the chair before sending."
    p.keywords = "AI-generated; Secure MOM; minutes of meeting; Medpark"
    p.comments = "Generated locally. AI Act Art. 50(2) marking."
    p.language = {"ro": "ro-MD", "ru": "ru-RU", "en": "en-GB"}[lang]


def _run(par, text, font, size, colour, bold=False):
    r = par.add_run(text)
    r.font.name, r.font.size, r.font.color.rgb, r.font.bold = font, Pt(size), colour, bold
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return r


def _para(doc, text, font, size, colour, bold=False, before=0, after=4):
    par = doc.add_paragraph()
    if text:
        _run(par, text, font, size, colour, bold)
    par.paragraph_format.space_before, par.paragraph_format.space_after = Pt(before), Pt(after)
    return par


def _item(doc, label, text, detail, newline=False):
    t = doc.add_table(rows=1, cols=2)
    a, b = t.rows[0].cells
    _widths(t, [Mm(25), Mm(143)])
    _cell(a, label, HEAD, 9, TEAL)
    par = b.paragraphs[0]
    _run(par, text, BODY, 10.5, CHARCOAL)
    if detail:
        _run(par, ("\n" if newline else "  ") + detail, BODY, 9, SLATE)


def _cell(cell, text, font, size, colour, new=False, bold=False):
    par = cell.add_paragraph() if new else cell.paragraphs[0]
    _run(par, text, font, size, colour, bold)
    par.paragraph_format.space_after = Pt(1)
    return par


def _table(doc, head, rows, widths):
    t = doc.add_table(rows=1, cols=len(head))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]
        _shade(c, "008286")
        _cell(c, h, HEAD, 9, RGBColor(0xFF, 0xFF, 0xFF), bold=True)
    for n, values in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(values):
            colour = GREY if values[i] in ("nestabilit", "не установлен", "not set") else CHARCOAL
            _cell(cells[i], v, BODY, 10, colour)
            if n % 2:
                _shade(cells[i], "F2F6F6")
    _widths(t, widths)


def _widths(table, widths):
    """Word ignores cell widths unless autofit is off and the grid agrees."""
    table.autofit = False
    tbl = table._tbl
    grid = tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if i < len(widths):
            gc.set(qn("w:w"), str(int(widths[i].twips)))
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w
    tblpr = tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:w"), str(int(sum(w.twips for w in widths))))
    tblw.set(qn("w:type"), "dxa")
    if tblw.getparent() is None:
        tblpr.append(tblw)


def _shade(cell, hex_colour):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_colour)
    cell._tc.get_or_add_tcPr().append(shd)


def _field(par, code):
    run = par.add_run()
    for tag, attr in (("w:fldChar", "begin"), ("w:instrText", code), ("w:fldChar", "end")):
        el = OxmlElement(tag)
        if tag == "w:fldChar":
            el.set(qn("w:fldCharType"), attr)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = attr
        run._element.append(el)
    run.font.size, run.font.name = Pt(7), HEAD
