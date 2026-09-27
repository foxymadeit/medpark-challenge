"""What the documents leave out: empty fields, talk time, unproven items, the
letterhead, and a second title or date. Checked in the PDF text and the DOCX,
in every language."""

import shutil
import subprocess

import pytest
from docx import Document

from mom.export import meeting_json
from mom.pipeline import attendees, render_documents
from mom.schemas import Fact, Line, Meeting, to_dict
from mom.write import fallback_body, plan, prompt

BODY = r"""\summary{S1}{The Board agreed to renew the MRI contract.}
\begin{agenda}
\agendaitem{T1}{MRI contract}
\agendaitem{T2}{Luminal pilot}
\end{agenda}
\topic{T1}{MRI contract}
\decision{D1}{The Board agreed to renew the contract.}{ }
\action{A1}{Participant 2}{}{Check the renewal terms.}
\action{A2}{}{}{Send the documents.}
\needsconfirmation{C1}{Unproven item from an older run.}
\topic{T2}{Luminal pilot}
\needsconfirmation{C2}{A topic with nothing proven under it.}
"""
# no number, place, chair or secretary: none of them may show up
BARE = Meeting(type="medical", date="2026-09-26", start="09:00", end="09:06")
TALK = {"ro": "12 min", "ru": "12 мин", "en": "12 min"}   # render files written before this change carry talk time
DOC = {"ro": "PROCES-VERBAL", "ru": "ПРОТОКОЛ", "en": "MINUTES"}
DATE = {"ro": "26.09.2026", "ru": "26.09.2026", "en": "26 September 2026"}
AI = {"ro": "Generat local de AI (test-model)", "ru": "Составлено локально ИИ (test-model)", "en": "Drafted locally by AI (test-model)"}
ABSENT = {  # labels of empty fields, the old placeholders and the letterhead
    "ro": ["Locul", "Președinte de ședință", "Secretar", "nestabilit", "Termen", "De confirmat", "Confidențial"],
    "ru": ["Место", "Председатель", "Секретарь", "не установлен", "Срок", "Требует подтверждения", "Конфиденциально"],
    "en": ["Place", "Chair", "Secretary", "not set", "Due", "Needs confirmation", "Confidential"],
}
OWNER = {"ro": "Responsabil:", "ru": "Ответственный:", "en": "Owner:"}


@pytest.fixture(scope="module")
def documents(tmp_path_factory):
    out = tmp_path_factory.mktemp("clean")
    state = {"meeting": to_dict(BARE), "bodies": {lang: BODY for lang in DOC},
             "attendees": {lang: [{"name": "Participant 1", "role": TALK[lang]}] for lang in DOC},
             "model": "test-model", "verified": "3/5"}
    files = render_documents(state, out)
    texts = {}
    for lang, f in files.items():
        pdf = subprocess.run(["pdftotext", f["pdf"], "-"], capture_output=True, text=True).stdout
        d = Document(f["docx"])
        sec = d.sections[0]
        parts = [p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells]
        parts += [p.text for p in sec.header.paragraphs + sec.footer.paragraphs]
        texts[lang] = {"pdf": pdf, "docx": "\n".join(parts)}
    return texts


def _each(documents):
    return [(lang, kind, text) for lang, t in documents.items() for kind, text in t.items()]


@pytest.mark.skipif(not shutil.which("xelatex"), reason="needs TeX Live")
def test_empty_fields_and_placeholders_are_left_out(documents):
    for lang, kind, text in _each(documents):
        for label in ABSENT[lang]:
            assert label not in text, (lang, kind, label)
        assert text.count(OWNER[lang]) == 1, (lang, kind)   # A1 has an owner, A2 has none


@pytest.mark.skipif(not shutil.which("xelatex"), reason="needs TeX Live")
def test_title_and_date_are_printed_once_and_no_number_is_invented(documents):
    for lang, kind, text in _each(documents):
        assert text.count(DOC[lang]) == 1 and text.count(DATE[lang]) == 1, (lang, kind)
        assert "09:00" in text and " 1\n" not in text.split(DOC[lang])[1][:40], (lang, kind)
        assert "09:06" not in text, (lang, kind)   # the start time only, never when the meeting ended


@pytest.mark.skipif(not shutil.which("xelatex"), reason="needs TeX Live")
def test_attendees_by_name_only_no_talk_time(documents):
    for lang, kind, text in _each(documents):
        assert "Participant 1" in text and TALK[lang] not in text, (lang, kind)


@pytest.mark.skipif(not shutil.which("xelatex"), reason="needs TeX Live")
def test_unproven_items_and_the_letterhead_are_gone_but_the_ai_notice_stays(documents):
    for lang, kind, text in _each(documents):
        assert "Unproven item" not in text and "Luminal pilot" not in text, (lang, kind)
        assert "Andrei Doga" not in text and "Medpark |" not in text and "+373" not in text, (lang, kind)
        assert AI[lang] in text.replace("\n", " "), (lang, kind)


def test_the_pipeline_lists_attendees_without_talk_time():
    people = attendees([Line("L0001", 0.0, 900.0, "Speaker 1", "Bună ziua.")], "ru")
    assert people == [{"name": "Участник 1"}]


FACTS = [
    Fact("T1", "topic", "Contract", ["L0001"]),
    Fact("D1", "decision", "Se aprobă.", ["L0001"], topic="T1"),
    Fact("A1", "action", "Sună furnizorul.", ["L0001"], topic="T1", owner="Speaker 9", status="confirm"),
    Fact("T2", "topic", "Pilot", ["L0002"], status="confirm"),
    Fact("N2", "note", "Pilotul începe.", ["L0002"], topic="T2", status="confirm"),
]


def test_unproven_facts_never_reach_the_writer_but_stay_in_the_app():
    rows = plan(FACTS, "ro")
    assert [r["id"] for r in rows] == ["S1", "T1", "D1"]
    assert "needsconfirmation" not in fallback_body(rows, "ro") and "needsconfirmation" not in prompt("ro")
    app = meeting_json(Meeting(type="medical", date="2026-09-26"), FACTS, [Line("L0001", 0, 5, "Speaker 1", "x")],
                       "\\topic{T1}{Contract}\n\\decision{D1}{Se aprobă.}{}\n", "ro")
    assert app["reviewFlags"] == ["Sună furnizorul.", "Pilot", "Pilotul începe."]
