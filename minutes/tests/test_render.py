import shutil
import subprocess
from pathlib import Path

import pytest
from docx import Document

from mom import latexcheck, render_docx, render_pdf
from mom.schemas import Meeting

BODY = r"""
\begin{agenda}
\agendaitem{T1}{Contractul RMN}
\end{agenda}
\topic{T1}{Contractul RMN}
\presented{N1}{dl Victor Munteanu}{Dl Victor Munteanu a comunicat că reducerea este de 10\%.}
\decision{D1}{Se aprobă reînnoirea contractului.}{în unanimitate}
\action{A1}{dl Victor Munteanu}{}{Verifică condițiile.}
\needsconfirmation{C1}{Cine trimite documentele.}
"""
MEETING = Meeting(type="medical", number="7", date="2026-09-26", start="14:00", end="14:30", place="Sala {1} & 50%",
                  chair=r"dna Ana POPESCU \input{x}", secretary="dna Elena Ciobanu",
                  attendees=[{"name": "dna Ana Popescu", "role": "director medical"}])


def test_docx_has_every_fact_the_ai_marking_and_private_permissions(tmp_path):
    out = render_docx.render(MEETING, latexcheck.parse(BODY), "ro", "test-model", "4/4", tmp_path / "m.docx")
    doc = Document(out)
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    for needle in ("PROCES-VERBAL nr. 7", "Se aprobă reînnoirea contractului.", "Verifică condițiile.", "nestabilit",
                   "Cine trimite documentele.", "reducerea este de 10%"):
        assert needle in text, needle
    assert "AI-generated" in doc.core_properties.keywords and "test-model" in doc.core_properties.author
    assert oct(out.stat().st_mode & 0o777) == "0o600"


def test_labels_come_from_the_latex_class_in_every_language():
    for lang, word in (("ro", "PROCES-VERBAL"), ("ru", "ПРОТОКОЛ"), ("en", "MINUTES")):
        s = render_docx.strings(lang)
        assert s["doc"] == word and s["owner"] and s["confirm"]


def test_header_values_are_escaped_before_they_reach_latex():
    tex = render_pdf.tex_source(MEETING, BODY, "ro", "m", "4/4")
    assert r"\input{x}" not in tex and r"Sala (1) \& 50\%" in tex


@pytest.mark.skipif(not shutil.which("latexmk"), reason="needs TeX Live")
def test_pdf_compiles_with_ai_marking_in_its_metadata(tmp_path):
    pdf = render_pdf.compile_pdf(render_pdf.tex_source(MEETING, BODY, "ro", "test-model", "4/4"),
                                 render_pdf.xmp_source(MEETING, "ro", "test-model"), tmp_path / "m.pdf")
    info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
    assert "AI-generated" in info and "test-model" in info
    text = subprocess.run(["pdftotext", str(pdf), "-"], capture_output=True, text=True).stdout
    assert "Se aprobă reînnoirea contractului." in text and "PROCES-VERBAL" in text
